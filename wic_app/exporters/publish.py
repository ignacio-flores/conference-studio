from __future__ import annotations

import json
import html
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import xlsxwriter

from engine.config import load_conference_config
from engine.scheduling import build_equal_time_ranges
from exporters.common import group_sessions_by_day_block, ordered_sessions, paper_row_lookup
from public_data import (
    PUBLIC_JSON_FILE,
    PUBLIC_XLSX_FILE,
    build_public_payload,
    resolve_public_export_settings,
    write_public_payload,
)
from reclassification_engine import (
    DAY_ORDER,
    DRAFT_OUTPUT_FILE,
    EXPORT_DIR,
    PUBLISH_DOCX_FILE,
    PUBLISH_PDF_FILE,
    PUBLISH_XLSX_FILE,
    ProgrammeState,
    format_minutes,
    parse_start_minutes,
    room_sort_key,
)

PLENARY_PROGRAMME_URL = "https://inequalitylab.world/www-site/uploads/2026/04/2026-WIC-Programme.pdf"
PLENARY_PROGRAMME_LINK_TEXT = "plenary sessions programme"
PARALLEL_SESSIONS_DISCLAIMER_PREFIX = "This programme covers parallel sessions only. For plenary sessions, see the "
PARALLEL_SESSIONS_DISCLAIMER = f"{PARALLEL_SESSIONS_DISCLAIMER_PREFIX}{PLENARY_PROGRAMME_LINK_TEXT}."
PROGRAMME_CHANGE_NOTICE = (
    "Programme details may still change before the conference, especially due to cancellations; "
    "sessions and presenters may be swapped."
)
PUBLISH_DISPLAY_FULL = "full"
PUBLISH_DISPLAY_PUBLIC_SAFE = "public_safe"
VALID_PUBLISH_DISPLAY_MODES = {PUBLISH_DISPLAY_FULL, PUBLISH_DISPLAY_PUBLIC_SAFE}


def normalize_publish_display_mode(value: Optional[str]) -> str:
    normalized = str(value or "").strip().lower()
    if normalized in VALID_PUBLISH_DISPLAY_MODES:
        return normalized
    return PUBLISH_DISPLAY_FULL


def _show_room_numbers(publish_display: str) -> bool:
    return normalize_publish_display_mode(publish_display) == PUBLISH_DISPLAY_FULL


def _show_moderator_labels(publish_display: str) -> bool:
    return normalize_publish_display_mode(publish_display) == PUBLISH_DISPLAY_FULL


def _show_change_notice(publish_display: str) -> bool:
    return normalize_publish_display_mode(publish_display) == PUBLISH_DISPLAY_PUBLIC_SAFE


def _display_room_value(room: object, publish_display: str) -> str:
    if not _show_room_numbers(publish_display):
        return ""
    return str(room or "").strip()


def _display_session_code_value(session_code: object, publish_display: str) -> str:
    if normalize_publish_display_mode(publish_display) == PUBLISH_DISPLAY_PUBLIC_SAFE:
        return ""
    return str(session_code or "").strip()


def _session_heading_text(session, publish_display: str) -> str:
    title = str(getattr(session, "session_title", "") or "").strip() or "[Untitled session]"
    session_code = _display_session_code_value(getattr(session, "session_code", ""), publish_display)
    return f"{session_code} | {title}" if session_code else title


def _session_slot_ranges(session) -> List[Tuple[int, int]]:
    capacity = max(1, int(getattr(session, "capacity", len(getattr(session, "papers", [])) or 1)))
    return build_equal_time_ranges(
        int(getattr(session, "start_min", 0) or 0),
        int(getattr(session, "end_min", 0) or 0),
        capacity,
    )


def _session_block_bounds(session) -> Tuple[int, int]:
    start_min = int(getattr(session, "start_min", parse_start_minutes(session.time)) or parse_start_minutes(session.time))
    end_min = int(getattr(session, "end_min", start_min + 90) or (start_min + 90))
    if end_min <= start_min:
        end_min = start_min + 90
    return start_min, end_min


def _build_agenda_cell(session, publish_display: str = PUBLISH_DISPLAY_FULL) -> str:
    if session is None:
        return ""
    lines = [_session_heading_text(session, publish_display)]
    slot_ranges = _session_slot_ranges(session)
    for idx, paper in enumerate(session.papers, start=1):
        if idx - 1 < len(slot_ranges):
            start_min, end_min = slot_ranges[idx - 1]
        else:
            start_min, end_min = _session_block_bounds(session)
        if paper is None:
            lines.append(f"{idx}. {format_minutes(start_min)}-{format_minutes(end_min)} [Reserve slot]")
        else:
            presenter = _presenter_label(paper, include_moderator=True)
            lines.append(f"{idx}. {format_minutes(start_min)}-{format_minutes(end_min)} {presenter} | {paper.title}")
    if session.overflow_papers:
        lines.append(f"Overflow: {len(session.overflow_papers)}")
    return "\n".join(lines)


def _presenter_label(paper, include_moderator: bool = False) -> str:
    presenter = str(getattr(paper, "full_name", "") or "").strip() or "[No presenter]"
    if include_moderator and bool(getattr(paper, "is_moderator", False)):
        return f"{presenter} (Chair)"
    return presenter


def _paper_title_and_presenter(paper, include_moderator: bool = False) -> Tuple[str, str]:
    if paper is None:
        return "[Reserve slot]", ""
    title = str(getattr(paper, "title", "") or "").strip() or "[Untitled]"
    presenter = _presenter_label(paper, include_moderator=include_moderator)
    return title, presenter


def _session_presentation_papers(session) -> List[object]:
    scheduled = list(getattr(session, "papers", []) or [])
    reserves = [paper for paper in scheduled if paper is None]
    presentations = [paper for paper in scheduled if paper is not None]
    presentations.extend(list(getattr(session, "overflow_papers", []) or []))
    presentations.extend(reserves)
    return presentations


def _publish_sessions_by_day(state: ProgrammeState) -> Tuple[List[object], List[str]]:
    ordered_sessions = sorted(
        state.sessions,
        key=lambda s: (s.day_num, parse_start_minutes(s.time), room_sort_key(s.room)),
    )
    return ordered_sessions, _ordered_day_labels_from_state(state)


def _write_title_presenter_cell(
    ws,
    row: int,
    col: int,
    paper,
    cell_fmt,
    presenter_fmt,
    include_moderator: bool = False,
) -> None:
    title, presenter = _paper_title_and_presenter(paper, include_moderator=include_moderator)
    if not presenter:
        ws.write(row, col, title, cell_fmt)
        return
    ws.write_rich_string(row, col, title, "\n", presenter_fmt, presenter, cell_fmt)


def _write_session_programme_cell(
    ws,
    row: int,
    col: int,
    session,
    cell_fmt,
    presenter_fmt,
    session_fmt,
    publish_display: str = PUBLISH_DISPLAY_FULL,
    include_moderator: bool = True,
) -> None:
    if session is None:
        ws.write(row, col, "", cell_fmt)
        return

    presentation_papers = _session_presentation_papers(session)
    if not presentation_papers:
        ws.write(row, col, _session_heading_text(session, publish_display), cell_fmt)
        return

    fragments: List[object] = [
        session_fmt,
        _session_heading_text(session, publish_display),
    ]
    for paper in presentation_papers:
        title, presenter = _paper_title_and_presenter(paper, include_moderator=include_moderator)
        fragments.extend(["\n\n", title])
        if presenter:
            fragments.extend(["\n", presenter_fmt, presenter])
    ws.write_rich_string(row, col, *fragments, cell_fmt)


def _ordered_day_labels_from_state(state: ProgrammeState) -> List[str]:
    configured_rank = {day: idx for idx, day in enumerate(DAY_ORDER)}
    seen_days = {session.day_label for session in state.sessions if session.day_label}
    if not seen_days:
        return list(DAY_ORDER)
    return sorted(seen_days, key=lambda day: (configured_rank.get(day, 999), day))


def _safe_sheet_name(raw: str, used: set[str]) -> str:
    base = str(raw or "Sheet").strip() or "Sheet"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate in used:
        suffix_text = f" ({suffix})"
        candidate = f"{base[: max(1, 31 - len(suffix_text))]}{suffix_text}"
        suffix += 1
    used.add(candidate)
    return candidate


def _write_day_sheet(
    ws,
    day_name: str,
    day_blocks: Dict[Tuple[int, str, str], Dict[str, object]],
    paper_row_lookup: Dict[str, int],
    fmts: Dict[str, object],
    include_comments: bool,
) -> None:
    rooms = sorted(
        {room for room_map in day_blocks.values() for room in room_map.keys()},
        key=room_sort_key,
    )

    ws.set_column(0, 0, 18)
    ws.set_column(1, max(1, len(rooms)), 56)
    ws.freeze_panes(2, 1)

    row = 0
    ws.merge_range(
        row,
        0,
        row,
        max(1, len(rooms)),
        f"{day_name} | presentation schedule",
        fmts["day_title"],
    )
    row += 1
    ws.write(row, 0, "Time Slot", fmts["header"])
    for col_idx, room in enumerate(rooms, start=1):
        ws.write(row, col_idx, room, fmts["header"])
    row += 1

    for block_key in sorted(day_blocks.keys(), key=lambda x: x[0]):
        _, time_label, block_label = block_key
        room_records = day_blocks[block_key]
        sessions_in_block = [session for session in room_records.values() if session is not None]
        if sessions_in_block:
            starts = [_session_block_bounds(session)[0] for session in sessions_in_block]
            ends = [_session_block_bounds(session)[1] for session in sessions_in_block]
            block_start = min(starts)
            block_end = max(ends)
        else:
            block_start = parse_start_minutes(time_label)
            block_end = block_start + 90
        ws.merge_range(
            row,
            0,
            row,
            max(1, len(rooms)),
            f"{block_label} | {format_minutes(block_start)}-{format_minutes(block_end)}",
            fmts["block"],
        )
        row += 1

        compatible_capacities = {
            max(1, int(getattr(session, "capacity", len(getattr(session, "papers", [])) or 1)))
            for session in sessions_in_block
        }
        compatible_bounds = {_session_block_bounds(session) for session in sessions_in_block}
        use_grid_mode = bool(sessions_in_block) and len(compatible_capacities) == 1 and len(compatible_bounds) == 1

        if use_grid_mode:
            capacity = next(iter(compatible_capacities))
            start_min, end_min = next(iter(compatible_bounds))
            slot_ranges = build_equal_time_ranges(start_min, end_min, capacity)
            for talk_idx in range(capacity):
                slot_start, slot_end = slot_ranges[talk_idx]
                ws.set_row(row, 74)
                ws.write(row, 0, f"{format_minutes(slot_start)}-{format_minutes(slot_end)}", fmts["time"])

                for col_idx, room in enumerate(rooms, start=1):
                    session = room_records.get(room)
                    if session is None:
                        ws.write(row, col_idx, "", fmts["cell"])
                        continue

                    paper = session.papers[talk_idx] if talk_idx < len(session.papers) else None
                    if paper is None:
                        ws.write(row, col_idx, "[Reserve slot]", fmts["reserve"])
                        continue

                    paper_row = paper_row_lookup[paper.submission_id] + 1
                    display = f"{paper.full_name}\n{paper.title}"
                    ws.write_url(
                        row,
                        col_idx,
                        f"internal:'Paper_Catalog'!A{paper_row}",
                        fmts["url_wrap"],
                        string=display,
                    )
                    if include_comments and paper.abstract:
                        ws.write_comment(row, col_idx, paper.abstract)
                row += 1
        else:
            ws.set_row(row, 210)
            ws.write(row, 0, "Mixed capacities: agenda view", fmts["time"])
            for col_idx, room in enumerate(rooms, start=1):
                session = room_records.get(room)
                ws.write(row, col_idx, _build_agenda_cell(session), fmts["wrap"])
            row += 1

        ws.write(row, 0, "Theme | Subtheme", fmts["small_note"])
        for col_idx, room in enumerate(rooms, start=1):
            session = room_records.get(room)
            if session is None:
                ws.write(row, col_idx, "", fmts["cell"])
                continue
            ws.write(
                row,
                col_idx,
                f"{session.primary_theme}\n{session.subtheme}",
                fmts["small_note_wrap"],
            )
        row += 2


def _write_publish_day_sheet(
    ws,
    day_name: str,
    day_blocks: Dict[Tuple[int, str, str], Dict[str, object]],
    fmts: Dict[str, object],
    publish_display: str = PUBLISH_DISPLAY_FULL,
) -> None:
    rooms = sorted(
        {room for room_map in day_blocks.values() for room in room_map.keys()},
        key=room_sort_key,
    )
    show_room_numbers = _show_room_numbers(publish_display)
    show_moderator_labels = _show_moderator_labels(publish_display)

    ws.set_column(0, 0, 18)
    ws.set_column(1, max(1, len(rooms)), 58)
    ws.freeze_panes(2, 1)

    row = 0
    ws.merge_range(
        row,
        0,
        row,
        max(1, len(rooms)),
        f"{day_name} | session schedule",
        fmts["day_title"],
    )
    row += 1
    ws.write(row, 0, "Session time", fmts["header"])
    for col_idx, room in enumerate(rooms, start=1):
        room_header = room if show_room_numbers else f"Track {col_idx}"
        ws.write(row, col_idx, room_header, fmts["header"])
    row += 1

    for block_key in sorted(day_blocks.keys(), key=lambda x: x[0]):
        _, time_label, block_label = block_key
        room_records = day_blocks[block_key]
        sessions_in_block = [session for session in room_records.values() if session is not None]
        if sessions_in_block:
            starts = [_session_block_bounds(session)[0] for session in sessions_in_block]
            ends = [_session_block_bounds(session)[1] for session in sessions_in_block]
            block_start = min(starts)
            block_end = max(ends)
            time_display = f"{format_minutes(block_start)}-{format_minutes(block_end)}"
        else:
            time_display = time_label
        ws.merge_range(
            row,
            0,
            row,
            max(1, len(rooms)),
            f"{block_label} | {time_display}",
            fmts["block"],
        )
        row += 1

        max_papers = max((len(_session_presentation_papers(session)) for session in sessions_in_block), default=1)
        ws.set_row(row, max(88, 34 + max_papers * 38))
        ws.write(row, 0, time_display, fmts["time"])
        for col_idx, room in enumerate(rooms, start=1):
            _write_session_programme_cell(
                ws,
                row,
                col_idx,
                room_records.get(room),
                fmts["wrap"],
                fmts["presenter_text"],
                fmts["session_text"],
                publish_display=publish_display,
                include_moderator=show_moderator_labels,
            )
        row += 2


def export_draft_workbook(state: ProgrammeState, output_path: Path = DRAFT_OUTPUT_FILE) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = xlsxwriter.Workbook(str(output_path))

    header_fmt = wb.add_format(
        {"bold": True, "bg_color": "#E6EEF7", "border": 1, "text_wrap": True, "valign": "top"}
    )
    cell_fmt = wb.add_format({"border": 1, "valign": "top"})
    wrap_fmt = wb.add_format({"border": 1, "valign": "top", "text_wrap": True})
    reserve_fmt = wb.add_format({"border": 1, "valign": "top", "italic": True, "font_color": "#7A4E00"})
    url_fmt = wb.add_format({"font_color": "blue", "underline": 1, "border": 1, "valign": "top"})
    url_wrap_fmt = wb.add_format(
        {"font_color": "blue", "underline": 1, "border": 1, "valign": "top", "text_wrap": True}
    )
    small_note_fmt = wb.add_format({"font_color": "#555555", "italic": True})
    small_note_wrap_fmt = wb.add_format({"font_color": "#555555", "italic": True, "text_wrap": True})
    block_fmt = wb.add_format(
        {"bold": True, "bg_color": "#DCE6F1", "border": 1, "align": "left", "valign": "vcenter"}
    )
    day_title_fmt = wb.add_format(
        {"bold": True, "bg_color": "#C9DAF8", "border": 1, "align": "left", "valign": "vcenter"}
    )
    time_fmt = wb.add_format({"bold": True, "bg_color": "#F4F7FC", "border": 1, "valign": "top"})

    fmts = {
        "header": header_fmt,
        "cell": cell_fmt,
        "wrap": wrap_fmt,
        "reserve": reserve_fmt,
        "url": url_fmt,
        "url_wrap": url_wrap_fmt,
        "small_note": small_note_fmt,
        "small_note_wrap": small_note_wrap_fmt,
        "block": block_fmt,
        "day_title": day_title_fmt,
        "time": time_fmt,
    }

    sorted_papers = sorted(state.papers, key=lambda p: p.submission_id)
    paper_row_lookup_map = paper_row_lookup(state)

    # Paper_Catalog
    ws_paper = wb.add_worksheet("Paper_Catalog")
    paper_headers = [
        "SubmissionID",
        "FullName",
        "EmailAddress",
        "ReviewerScore",
        "SourceThemes",
        "PrimaryTheme",
        "SecondaryThemeTags",
        "DetailedSubtheme",
        "Reviewed",
        "OverrideNotes",
        "SessionCode",
        "SessionTitle",
        "AssignmentRationale",
        "Title",
        "Abstract",
        "LinkToPDF",
        "Day",
        "Time",
        "Block",
        "Room",
    ]
    for c, h in enumerate(paper_headers):
        ws_paper.write(0, c, h, header_fmt)

    ws_paper.freeze_panes(1, 0)
    ws_paper.set_column(0, 0, 14)
    ws_paper.set_column(1, 2, 24)
    ws_paper.set_column(3, 3, 12)
    ws_paper.set_column(4, 7, 34)
    ws_paper.set_column(8, 12, 28)
    ws_paper.set_column(13, 13, 46)
    ws_paper.set_column(14, 14, 78)
    ws_paper.set_column(15, 15, 52)
    ws_paper.set_column(16, 19, 18)

    for row_idx, paper in enumerate(sorted_papers, start=1):
        ws_paper.write(row_idx, 0, paper.submission_id, cell_fmt)
        ws_paper.write(row_idx, 1, paper.full_name, cell_fmt)
        ws_paper.write(row_idx, 2, paper.email, cell_fmt)
        ws_paper.write(row_idx, 3, paper.reviewer_score, cell_fmt)
        ws_paper.write(row_idx, 4, paper.source_themes, wrap_fmt)
        ws_paper.write(row_idx, 5, paper.primary_theme, wrap_fmt)
        ws_paper.write(row_idx, 6, ", ".join(paper.secondary_tags), wrap_fmt)
        ws_paper.write(row_idx, 7, paper.detailed_subtheme, wrap_fmt)
        ws_paper.write(row_idx, 8, "True" if paper.reviewed else "False", cell_fmt)
        ws_paper.write(row_idx, 9, paper.override_notes, wrap_fmt)
        ws_paper.write(row_idx, 10, paper.session_code, cell_fmt)
        ws_paper.write(row_idx, 11, paper.session_title, wrap_fmt)
        ws_paper.write(row_idx, 12, paper.rationale, wrap_fmt)
        ws_paper.write(row_idx, 13, paper.title, wrap_fmt)
        ws_paper.write(row_idx, 14, paper.abstract, wrap_fmt)
        if paper.link_to_pdf.startswith("http"):
            ws_paper.write_url(row_idx, 15, paper.link_to_pdf, url_fmt, string=paper.link_to_pdf)
        else:
            ws_paper.write(row_idx, 15, paper.link_to_pdf, cell_fmt)
        ws_paper.write(row_idx, 16, paper.day_label, cell_fmt)
        ws_paper.write(row_idx, 17, paper.time, cell_fmt)
        ws_paper.write(row_idx, 18, paper.block_label, cell_fmt)
        ws_paper.write(row_idx, 19, paper.room, cell_fmt)

    # Session_Assignments
    ws_session = wb.add_worksheet("Session_Assignments")
    session_rows = ordered_sessions(state)
    max_capacity = max([max(1, int(getattr(session, "capacity", len(session.papers) or 1))) for session in session_rows], default=1)
    session_headers = [
        "SessionCode",
        "Day",
        "Time",
        "Block",
        "Room",
        "Capacity",
        "SessionTitle",
        "PrimaryTheme",
        "Subtheme",
    ]
    for talk_idx in range(1, max_capacity + 1):
        session_headers.extend(
            [
                f"Slot{talk_idx}_SubmissionID",
                f"Slot{talk_idx}_Presenter",
                f"Slot{talk_idx}_Title",
            ]
        )
    for c, h in enumerate(session_headers):
        ws_session.write(0, c, h, header_fmt)

    ws_session.freeze_panes(1, 0)
    ws_session.set_column(0, 0, 16)
    ws_session.set_column(1, 4, 18)
    ws_session.set_column(5, 8, 42)
    base_slot_col = 9
    for talk_idx in range(max_capacity):
        offset = base_slot_col + talk_idx * 3
        ws_session.set_column(offset, offset, 14)
        ws_session.set_column(offset + 1, offset + 1, 24)
        ws_session.set_column(offset + 2, offset + 2, 44)

    for row_idx, session in enumerate(session_rows, start=1):
        ws_session.set_row(row_idx, 70)
        ws_session.write(row_idx, 0, session.session_code, cell_fmt)
        ws_session.write(row_idx, 1, session.day_label, cell_fmt)
        ws_session.write(row_idx, 2, session.time, cell_fmt)
        ws_session.write(row_idx, 3, session.block_label, cell_fmt)
        ws_session.write(row_idx, 4, session.room, cell_fmt)
        ws_session.write(row_idx, 5, max(1, int(getattr(session, "capacity", len(session.papers) or 1))), cell_fmt)
        ws_session.write(row_idx, 6, session.session_title, wrap_fmt)
        ws_session.write(row_idx, 7, session.primary_theme, wrap_fmt)
        ws_session.write(row_idx, 8, session.subtheme, wrap_fmt)

        for talk_idx in range(max_capacity):
            paper = session.papers[talk_idx] if talk_idx < len(session.papers) else None
            base_col = 9 + talk_idx * 3
            if paper is None:
                ws_session.write(row_idx, base_col, "RESERVE_SLOT", reserve_fmt)
                ws_session.write(row_idx, base_col + 1, "[Reserve slot]", reserve_fmt)
                ws_session.write(row_idx, base_col + 2, "[Reserve slot]", reserve_fmt)
            else:
                ws_session.write(row_idx, base_col, paper.submission_id, cell_fmt)
                ws_session.write(row_idx, base_col + 1, paper.full_name, wrap_fmt)
                paper_row = paper_row_lookup_map[paper.submission_id] + 1
                ws_session.write_url(
                    row_idx,
                    base_col + 2,
                    f"internal:'Paper_Catalog'!A{paper_row}",
                    url_wrap_fmt,
                    string=paper.title,
                )
                if paper.abstract:
                    ws_session.write_comment(row_idx, base_col + 2, paper.abstract)

    # Day sheets
    day_blocks = group_sessions_by_day_block(state)
    ordered_days = _ordered_day_labels_from_state(state)
    day_sheet_names = {
        "Day 1 (4th June)": "Day 1 Programme",
        "Day 2 (5th June)": "Day 2 Programme",
        "Day 3 (6th June)": "Day 3 Programme",
    }
    used_sheet_names = {"Paper_Catalog", "Session_Assignments", "Reserves", "Themes_Themed"}
    for day_idx, day_name in enumerate(ordered_days, start=1):
        default_name = day_sheet_names.get(day_name, f"Day {day_idx} Programme")
        ws_day = wb.add_worksheet(_safe_sheet_name(default_name, used_sheet_names))
        _write_day_sheet(
            ws_day,
            day_name,
            day_blocks.get(day_name, {}),
            paper_row_lookup_map,
            fmts,
            include_comments=True,
        )

    # Reserves sheet
    ws_res = wb.add_worksheet("Reserves")
    reserve_headers = [
        "ReserveSlotID",
        "SessionCode",
        "Day",
        "Time",
        "Room",
        "SessionTitle",
        "PrimaryTheme",
        "Subtheme",
        "SuggestedAlternates",
    ]
    for c, h in enumerate(reserve_headers):
        ws_res.write(0, c, h, header_fmt)

    ws_res.freeze_panes(1, 0)
    ws_res.set_column(0, 0, 14)
    ws_res.set_column(1, 1, 16)
    ws_res.set_column(2, 4, 18)
    ws_res.set_column(5, 7, 44)
    ws_res.set_column(8, 8, 86)

    reserve_row = 1
    for session in session_rows:
        for talk_idx, paper in enumerate(session.papers, start=1):
            if paper is not None:
                continue

            candidates = [
                p
                for p in state.papers
                if p.session_code != session.session_code and p.primary_theme == session.primary_theme
            ]
            candidates = sorted(
                candidates,
                key=lambda p: (
                    1 if p.detailed_subtheme == session.subtheme else 0,
                    p.title.lower(),
                ),
                reverse=True,
            )[:5]
            candidate_text = "\n".join(
                f"{p.submission_id} - {p.full_name} - {p.title} ({p.session_code})" for p in candidates
            )

            ws_res.write(reserve_row, 0, f"RSV-{reserve_row:02d}-Slot{talk_idx}", cell_fmt)
            ws_res.write(reserve_row, 1, session.session_code, cell_fmt)
            ws_res.write(reserve_row, 2, session.day_label, cell_fmt)
            ws_res.write(reserve_row, 3, session.time, cell_fmt)
            ws_res.write(reserve_row, 4, session.room, cell_fmt)
            ws_res.write(reserve_row, 5, session.session_title, wrap_fmt)
            ws_res.write(reserve_row, 6, session.primary_theme, wrap_fmt)
            ws_res.write(reserve_row, 7, session.subtheme, wrap_fmt)
            ws_res.write(reserve_row, 8, candidate_text, wrap_fmt)
            reserve_row += 1

    # Themes summary
    ws_theme = wb.add_worksheet("Themes_Themed")
    ws_theme.write_row(0, 0, ["Theme", "Subtheme", "Sessions", "Papers"], header_fmt)
    ws_theme.freeze_panes(1, 0)
    ws_theme.set_column(0, 1, 44)
    ws_theme.set_column(2, 3, 12)

    summary = defaultdict(lambda: {"sessions": 0, "papers": 0})
    for session in session_rows:
        key = (session.primary_theme, session.subtheme)
        summary[key]["sessions"] += 1
        summary[key]["papers"] += sum(1 for p in session.papers if p is not None)

    row = 1
    for (theme, subtheme), stats in sorted(summary.items(), key=lambda x: (x[0][0], x[0][1])):
        ws_theme.write(row, 0, theme, wrap_fmt)
        ws_theme.write(row, 1, subtheme, wrap_fmt)
        ws_theme.write(row, 2, stats["sessions"], cell_fmt)
        ws_theme.write(row, 3, stats["papers"], cell_fmt)
        row += 1

    wb.close()
    return output_path


def export_public_payload(
    state: ProgrammeState,
    output_path: Path = PUBLIC_JSON_FILE,
    *,
    conference_config=None,
    publish_display: str = PUBLISH_DISPLAY_FULL,
    show_rooms: Optional[bool] = None,
    show_moderators: Optional[bool] = None,
    show_links: Optional[bool] = None,
) -> Path:
    conference = conference_config or load_conference_config()
    payload = build_public_payload(
        state,
        conference,
        publish_display=publish_display,
        show_rooms=show_rooms,
        show_moderators=show_moderators,
        show_links=show_links,
    )
    return write_public_payload(payload, output_path)


def export_public_excel(
    state: ProgrammeState,
    output_path: Path = PUBLIC_XLSX_FILE,
    *,
    conference_config=None,
    publish_display: str = PUBLISH_DISPLAY_FULL,
    show_rooms: Optional[bool] = None,
    show_moderators: Optional[bool] = None,
    show_links: Optional[bool] = None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    conference = conference_config or load_conference_config()
    payload = build_public_payload(
        state,
        conference,
        publish_display=publish_display,
        show_rooms=show_rooms,
        show_moderators=show_moderators,
        show_links=show_links,
    )
    public_settings = resolve_public_export_settings(
        public_settings=payload.get("public_settings"),
        publish_display=publish_display,
        show_rooms=show_rooms,
        show_moderators=show_moderators,
        show_links=show_links,
    )

    wb = xlsxwriter.Workbook(str(output_path))
    title_fmt = wb.add_format({"bold": True, "font_size": 20})
    subtitle_fmt = wb.add_format({"bold": True, "font_size": 14})
    header_fmt = wb.add_format({"bold": True, "bg_color": "#E8EEF8", "border": 1, "valign": "top"})
    cell_fmt = wb.add_format({"border": 1, "valign": "top"})
    wrap_fmt = wb.add_format({"border": 1, "valign": "top", "text_wrap": True})

    ws_cover = wb.add_worksheet("Cover")
    ws_cover.set_column(0, 0, 90)
    ws_cover.write(0, 0, payload["conference"]["title"], title_fmt)
    ws_cover.write(2, 0, payload["conference"]["subtitle"], subtitle_fmt)
    ws_cover.write(4, 0, f"Generated: {payload['conference']['generated_at']}", cell_fmt)
    ws_cover.write(6, 0, "This workbook is a public read-only programme export.", cell_fmt)
    ws_cover.write(8, 0, f"Sessions: {len(payload['sessions'])}", cell_fmt)
    ws_cover.write(9, 0, f"Papers: {len(payload['papers'])}", cell_fmt)
    ws_cover.write(
        11,
        0,
        "Public settings: "
        f"rooms={'shown' if public_settings['show_rooms'] else 'hidden'}, "
        f"moderators={'shown' if public_settings['show_moderators'] else 'hidden'}, "
        f"links={'shown' if public_settings['show_links'] else 'hidden'}",
        cell_fmt,
    )

    ws_sessions = wb.add_worksheet("Sessions")
    session_headers = [
        "SessionCode",
        "SessionTitle",
        "Day",
        "Time",
        "Block",
        "Room",
        "PrimaryTheme",
        "Subtheme",
        "Talks",
    ]
    ws_sessions.write_row(0, 0, session_headers, header_fmt)
    ws_sessions.freeze_panes(1, 0)
    ws_sessions.set_column(0, 0, 16)
    ws_sessions.set_column(1, 1, 40)
    ws_sessions.set_column(2, 5, 18)
    ws_sessions.set_column(6, 7, 32)
    ws_sessions.set_column(8, 8, 60)
    for row_idx, session in enumerate(payload["sessions"], start=1):
        ws_sessions.write(row_idx, 0, session["session_code"], cell_fmt)
        ws_sessions.write(row_idx, 1, session["session_title"], wrap_fmt)
        ws_sessions.write(row_idx, 2, session["day_label"], cell_fmt)
        ws_sessions.write(row_idx, 3, session["time"], cell_fmt)
        ws_sessions.write(row_idx, 4, session["block_label"], wrap_fmt)
        ws_sessions.write(row_idx, 5, session.get("display_room", session["room"]), cell_fmt)
        ws_sessions.write(row_idx, 6, session["primary_theme"], wrap_fmt)
        ws_sessions.write(row_idx, 7, session["subtheme"], wrap_fmt)
        ws_sessions.write(
            row_idx,
            8,
            "\n".join(
                f"{talk.get('display_presenter', talk.get('presenter_display', talk['authors']))} - {talk['title']}"
                for talk in session["talks"]
            ),
            wrap_fmt,
        )

    ws_papers = wb.add_worksheet("Papers")
    paper_headers = [
        "SubmissionID",
        "Presenter",
        "Title",
        "Abstract",
        "PaperURL",
        "SessionCode",
        "SessionTitle",
        "Day",
        "Time",
        "Room",
        "PrimaryTheme",
        "Subtheme",
    ]
    ws_papers.write_row(0, 0, paper_headers, header_fmt)
    ws_papers.freeze_panes(1, 0)
    ws_papers.set_column(0, 1, 18)
    ws_papers.set_column(2, 2, 52)
    ws_papers.set_column(3, 3, 80)
    ws_papers.set_column(4, 6, 24)
    ws_papers.set_column(7, 11, 18)
    for row_idx, paper in enumerate(payload["papers"], start=1):
        ws_papers.write(row_idx, 0, paper["submission_id"], cell_fmt)
        ws_papers.write(row_idx, 1, paper.get("display_presenter", paper.get("presenter_display", paper["authors"])), wrap_fmt)
        ws_papers.write(row_idx, 2, paper["title"], wrap_fmt)
        ws_papers.write(row_idx, 3, paper["abstract"], wrap_fmt)
        if paper.get("paper_url"):
            ws_papers.write_url(row_idx, 4, paper["paper_url"], cell_fmt, string=paper["paper_url"])
        else:
            ws_papers.write(row_idx, 4, "", cell_fmt)
        ws_papers.write(row_idx, 5, paper["session_code"], cell_fmt)
        ws_papers.write(row_idx, 6, paper["session_title"], wrap_fmt)
        ws_papers.write(row_idx, 7, paper["day_label"], cell_fmt)
        ws_papers.write(row_idx, 8, paper["time"], cell_fmt)
        ws_papers.write(row_idx, 9, paper.get("display_room", paper["room"]), cell_fmt)
        ws_papers.write(row_idx, 10, paper["primary_theme"], wrap_fmt)
        ws_papers.write(row_idx, 11, paper["subtheme"], wrap_fmt)

    wb.close()
    return output_path


def export_publish_excel(
    state: ProgrammeState,
    output_path: Path = PUBLISH_XLSX_FILE,
    publish_display: str = PUBLISH_DISPLAY_FULL,
    *,
    show_links: bool = False,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = xlsxwriter.Workbook(str(output_path))
    conference = load_conference_config()
    publish_display = normalize_publish_display_mode(publish_display)
    show_moderator_labels = _show_moderator_labels(publish_display)

    header_fmt = wb.add_format({"bold": True, "bg_color": "#E8EEF8", "border": 1, "valign": "top", "text_wrap": True})
    cell_fmt = wb.add_format({"border": 1, "valign": "top", "text_wrap": True})
    wrap_fmt = wb.add_format({"border": 1, "valign": "top", "text_wrap": True})
    day_title_fmt = wb.add_format(
        {"bold": True, "bg_color": "#C9DAF8", "border": 1, "align": "left", "valign": "vcenter"}
    )
    block_fmt = wb.add_format(
        {"bold": True, "bg_color": "#DFE8F7", "border": 1, "align": "left", "valign": "vcenter"}
    )
    time_fmt = wb.add_format({"bold": True, "bg_color": "#F4F7FC", "border": 1, "text_wrap": True})
    reserve_fmt = wb.add_format({"border": 1, "italic": True, "font_color": "#7A4E00", "text_wrap": True, "valign": "top"})
    presenter_text_fmt = wb.add_format({"font_color": "#666666", "italic": True})
    session_text_fmt = wb.add_format({"bold": True})
    link_fmt = wb.add_format({"font_color": "blue", "underline": 1, "text_wrap": True, "valign": "top"})
    cover_note_fmt = wb.add_format({"italic": True, "font_color": "#555555", "text_wrap": True})

    fmts = {
        "header": header_fmt,
        "cell": cell_fmt,
        "wrap": wrap_fmt,
        "reserve": reserve_fmt,
        "small_note": wrap_fmt,
        "small_note_wrap": wrap_fmt,
        "block": block_fmt,
        "day_title": day_title_fmt,
        "time": time_fmt,
        "presenter_text": presenter_text_fmt,
        "session_text": session_text_fmt,
    }

    ws_cover = wb.add_worksheet("Cover")
    ws_cover.set_column(0, 0, 90)
    ws_cover.write(0, 0, conference.conference_title, wb.add_format({"bold": True, "font_size": 20}))
    ws_cover.write(2, 0, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ws_cover.write_url(
        4,
        0,
        PLENARY_PROGRAMME_URL,
        link_fmt,
        string=PARALLEL_SESSIONS_DISCLAIMER,
    )
    stats_row_start = 6
    if _show_change_notice(publish_display):
        ws_cover.write(5, 0, PROGRAMME_CHANGE_NOTICE, cover_note_fmt)
        stats_row_start = 7
    ws_cover.write(stats_row_start, 0, f"Scheduled papers: {state.validations.get('scheduled_papers', 0)}")
    ws_cover.write(stats_row_start + 1, 0, f"Overflow papers: {state.validations.get('overflow_papers', 0)}")
    ws_cover.write(stats_row_start + 2, 0, f"Unassigned papers: {state.validations.get('unassigned_papers', 0)}")

    day_blocks = group_sessions_by_day_block(state)
    ordered_days = _ordered_day_labels_from_state(state)

    used_sheet_names = {"Cover"}
    for day_idx, day_name in enumerate(ordered_days, start=1):
        ws_day = wb.add_worksheet(_safe_sheet_name(f"Day {day_idx}", used_sheet_names))
        _write_publish_day_sheet(
            ws_day,
            day_name,
            day_blocks.get(day_name, {}),
            fmts,
            publish_display=publish_display,
        )

    wb.close()
    return output_path


def _load_branding(config_path: Optional[Path]) -> Dict[str, str]:
    defaults = {
        "conference_title": "World Inequality Conference 2026",
        "conference_subtitle": "Publish Programme",
        "primary_color": "#1F3C73",
    }
    if not config_path or not config_path.exists():
        return defaults
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except Exception:
        return defaults

    for key in defaults.keys():
        if key in data and isinstance(data[key], str) and data[key].strip():
            defaults[key] = data[key].strip()
    return defaults


def export_publish_pdf(
    state: ProgrammeState,
    output_path: Path = PUBLISH_PDF_FILE,
    branding_config_path: Optional[Path] = Path(__file__).resolve().parents[1] / "assets/branding.json",
    publish_display: str = PUBLISH_DISPLAY_FULL,
) -> Path:
    try:
        from runtime_compat import install_hashlib_usedforsecurity_compat

        install_hashlib_usedforsecurity_compat()
    except Exception:
        pass

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "PDF export requires `reportlab`. Install dependencies with `pip install -r wic_app/requirements.txt`."
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    publish_display = normalize_publish_display_mode(publish_display)
    show_room_numbers = _show_room_numbers(publish_display)
    show_moderator_labels = _show_moderator_labels(publish_display)

    branding = _load_branding(branding_config_path)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        fontSize=24,
        leading=30,
        textColor=colors.HexColor(branding["primary_color"]),
        spaceAfter=8,
    )
    subtitle_style = ParagraphStyle(
        "SubtitleStyle",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        textColor=colors.black,
        spaceAfter=8,
    )
    section_style = ParagraphStyle(
        "SectionStyle",
        parent=styles["Heading2"],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor(branding["primary_color"]),
        spaceBefore=10,
        spaceAfter=6,
    )
    small_style = ParagraphStyle(
        "SmallStyle",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        spaceAfter=3,
    )

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=14 * mm,
        title=branding["conference_title"],
    )

    story = []
    story.append(Paragraph(branding["conference_title"], title_style))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]))
    story.append(
        Paragraph(
            f"{html.escape(PARALLEL_SESSIONS_DISCLAIMER_PREFIX)}"
            f"<a href='{html.escape(PLENARY_PROGRAMME_URL)}'>{html.escape(PLENARY_PROGRAMME_LINK_TEXT)}</a>.",
            styles["Normal"],
        )
    )
    if _show_change_notice(publish_display):
        story.append(Paragraph(html.escape(PROGRAMME_CHANGE_NOTICE), styles["Normal"]))
    story.append(Spacer(1, 6 * mm))

    ordered_sessions, ordered_days = _publish_sessions_by_day(state)

    story.append(Paragraph("Session Booklet", section_style))

    for day_name in ordered_days:
        story.append(Paragraph(day_name, subtitle_style))
        day_sessions = [s for s in ordered_sessions if s.day_label == day_name]

        for session in day_sessions:
            session_heading = f"<b>{html.escape(session.session_title)}</b> | {html.escape(session.time)}"
            if show_room_numbers:
                session_heading = f"{session_heading} | {html.escape(session.room)}"
            story.append(
                Paragraph(
                    session_heading,
                    styles["Normal"],
                )
            )

            for idx, paper in enumerate(_session_presentation_papers(session), start=1):
                if paper is None:
                    story.append(Paragraph(f"{idx}. [Reserve slot]", small_style))
                    continue
                title, presenter = _paper_title_and_presenter(paper, include_moderator=show_moderator_labels)
                story.append(
                    Paragraph(
                        f"{idx}. {html.escape(title)} "
                        f"<font color='#666666'><i>{html.escape(presenter)}</i></font>",
                        small_style,
                    )
                )

            story.append(Spacer(1, 3 * mm))

        story.append(PageBreak())

    doc.build(story)
    return output_path


def _add_docx_hyperlink(paragraph, url: str, text: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def export_publish_docx(
    state: ProgrammeState,
    output_path: Path = PUBLISH_DOCX_FILE,
    branding_config_path: Optional[Path] = Path(__file__).resolve().parents[1] / "assets/branding.json",
    publish_display: str = PUBLISH_DISPLAY_FULL,
) -> Path:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Word export requires `python-docx`. Install dependencies with `pip install -r wic_app/requirements.txt`."
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    publish_display = normalize_publish_display_mode(publish_display)
    show_room_numbers = _show_room_numbers(publish_display)
    show_moderator_labels = _show_moderator_labels(publish_display)

    branding = _load_branding(branding_config_path)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)

    document.add_paragraph(branding["conference_title"], style="Title")
    document.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    disclaimer = document.add_paragraph()
    disclaimer.add_run(PARALLEL_SESSIONS_DISCLAIMER_PREFIX)
    _add_docx_hyperlink(disclaimer, PLENARY_PROGRAMME_URL, PLENARY_PROGRAMME_LINK_TEXT)
    disclaimer.add_run(".")
    if _show_change_notice(publish_display):
        document.add_paragraph(PROGRAMME_CHANGE_NOTICE)

    document.add_paragraph("Session Booklet", style="Heading 1")
    ordered_sessions, ordered_days = _publish_sessions_by_day(state)

    for day_index, day_name in enumerate(ordered_days):
        document.add_paragraph(day_name, style="Heading 2")
        day_sessions = [session for session in ordered_sessions if session.day_label == day_name]

        for session in day_sessions:
            session_heading = document.add_paragraph()
            title_run = session_heading.add_run(str(getattr(session, "session_title", "") or ""))
            title_run.bold = True
            session_heading.add_run(f" | {getattr(session, 'time', '')}")
            if show_room_numbers:
                session_heading.add_run(f" | {getattr(session, 'room', '')}")

            for idx, paper in enumerate(_session_presentation_papers(session), start=1):
                line = document.add_paragraph()
                if paper is None:
                    line.add_run(f"{idx}. [Reserve slot]")
                    continue
                title, presenter = _paper_title_and_presenter(paper, include_moderator=show_moderator_labels)
                line.add_run(f"{idx}. {title} ")
                presenter_run = line.add_run(presenter)
                presenter_run.italic = True

        if day_index < len(ordered_days) - 1:
            document.add_page_break()

    document.save(str(output_path))
    return output_path


def export_all(
    state: ProgrammeState,
    draft_output: Path = DRAFT_OUTPUT_FILE,
    publish_xlsx_output: Path = PUBLISH_XLSX_FILE,
    publish_pdf_output: Path = PUBLISH_PDF_FILE,
    publish_docx_output: Path = PUBLISH_DOCX_FILE,
    publish_display: str = PUBLISH_DISPLAY_FULL,
) -> Dict[str, Path]:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    return {
        "draft_xlsx": export_draft_workbook(state, draft_output),
        "publish_xlsx": export_publish_excel(state, publish_xlsx_output, publish_display=publish_display),
        "publish_pdf": export_publish_pdf(state, publish_pdf_output, publish_display=publish_display),
        "publish_docx": export_publish_docx(state, publish_docx_output, publish_display=publish_display),
    }
